"""Importadores. Hoje: planilha Compromissos_Financeiros.xlsx.
Próximos: faturas Itaú (.xlsx) e Nubank (.csv) com a regra "só titular".

Funcionam com Session (app) ou Connection (script)."""
import hashlib
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
from sqlalchemy import text

from core.semente import garantir_categoria

TIPO_POR_FORMA = {"Cartao de Credito": "cartao_credito", "Debito em Conta": "conta_corrente"}
COLUNAS = ["Descricao", "Categoria", "Subcategoria", "Forma de Pagamento",
           "Instrumento", "Emissor", "Valor", "Vencimento"]


def ler_planilha_compromissos(arquivo) -> list[dict]:
    """arquivo = caminho ou arquivo aberto (st.file_uploader)."""
    ws = openpyxl.load_workbook(arquivo, data_only=True)["Compromissos"]
    cab = None
    for row in ws.iter_rows(min_row=1, max_row=15, values_only=True):
        if "Descricao" in row:
            cab = list(row)
            break
    if cab is None:
        raise ValueError("Não achei o cabeçalho (Descricao, Categoria, …) na aba Compromissos.")
    idx = {n: cab.index(n) for n in COLUNAS}
    linhas = []
    for row in ws.iter_rows(min_row=1, values_only=True):
        desc, venc, valor = row[idx["Descricao"]], row[idx["Vencimento"]], row[idx["Valor"]]
        if not desc or desc == "Descricao" or not hasattr(venc, "year") or valor is None:
            continue                                  # cabeçalho, vazias e linha TOTAL
        linhas.append({k: row[i] for k, i in idx.items()})
    return linhas


def importar_planilha_compromissos(cx, dono_id, arquivo, nome_arquivo=None) -> dict:
    linhas = ler_planilha_compromissos(arquivo)

    # contas: uma por instrumento; dia de vencimento do cartão = o mais comum na planilha
    dias, info = defaultdict(Counter), {}
    for l in linhas:
        if l["Instrumento"]:
            dias[l["Instrumento"]][l["Vencimento"].day] += 1
            info[l["Instrumento"]] = (l["Emissor"], l["Forma de Pagamento"])
    contas = {}
    for nome, (emissor, forma) in info.items():
        tipo = TIPO_POR_FORMA.get(forma, "outro")
        dia = dias[nome].most_common(1)[0][0] if tipo == "cartao_credito" else None
        contas[nome] = cx.execute(text("""
            insert into contas (dono_id, nome, tipo, instituicao, dia_vencimento)
            values (:u, :n, cast(:t as tipo_conta), :i, :d)
            on conflict (dono_id, nome) do update set instituicao = coalesce(contas.instituicao, excluded.instituicao)
            returning id"""), dict(u=dono_id, n=nome, t=tipo, i=emissor, d=dia)).scalar()

    cats = {}
    for l in linhas:
        chave = (l["Categoria"], l["Subcategoria"])
        if chave not in cats:
            pai = garantir_categoria(cx, dono_id, l["Categoria"])
            cats[chave] = garantir_categoria(cx, dono_id, l["Subcategoria"], pai) if l["Subcategoria"] else pai

    vistos, novos = Counter(), 0
    for l in linhas:
        base = f"planilha|{l['Descricao']}|{float(l['Valor']):.2f}|{l['Vencimento']:%Y-%m-%d}|{l['Instrumento']}"
        vistos[base] += 1                              # mesma compra 2x no mesmo dia -> chaves diferentes
        chave = hashlib.sha1(f"{base}|{vistos[base]}".encode()).hexdigest()
        r = cx.execute(text("""
            insert into lancamentos (dono_id, conta_id, categoria_id, tipo, descricao, valor, data_vencimento,
                                     forma_pagamento, status, origem, chave_importacao)
            values (:u, :conta, :cat, 'despesa', :desc, :valor, :venc, :forma, 'previsto', 'migracao', :chave)
            on conflict (dono_id, chave_importacao) where chave_importacao is not null do nothing"""),
            dict(u=dono_id, conta=contas.get(l["Instrumento"]), cat=cats[(l["Categoria"], l["Subcategoria"])],
                 desc=l["Descricao"], valor=round(float(l["Valor"]), 2), venc=l["Vencimento"].date(),
                 forma=l["Forma de Pagamento"], chave=chave))
        novos += r.rowcount

    cx.execute(text("""insert into importacoes (dono_id, arquivo, formato, linhas_lidas, linhas_importadas)
                       values (:u, :a, 'planilha_compromissos', :l, :n)"""),
               dict(u=dono_id, a=nome_arquivo or Path(str(arquivo)).name, l=len(linhas), n=novos))
    return {"lidas": len(linhas), "novas": novos, "contas": len(contas), "categorias": len(cats),
            "total": round(sum(float(l["Valor"]) for l in linhas), 2)}
